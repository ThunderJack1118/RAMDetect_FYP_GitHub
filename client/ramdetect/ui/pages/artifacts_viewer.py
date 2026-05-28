"""
Artifacts Viewer — Page 7  (v7 — extraction-tool aligned)
──────────────────────────────────────────────────────────
Complete field-name audit against ramdetect_extraction_tool.py

PROCESS dict (from eprocess_scanner):
  pid          → p['pid']            integer
  ppid         → p['ppid']           integer
  image_name   → p['image_name']     ← primary name key  (NOT 'name')
  path         → p.get('path','')
  anomalies    → p.get('anomalies',[])
  status       → NOT reliably stored; compute via determine_status()
  is_suspicious→ p.get('is_suspicious')
  (NO thread_count / handle_count — must cross-ref from threads/handles lists)

THREAD dict (from thread_parser):
  tid          → t.get('tid', 0)
  pid          → t.get('pid', 0)
  (no process name — look up from pid→name map)
  start_address→ t.get('start_address', 0)   ← exact key
  state_name   → t.get('state_name', ...)    ← primary key (NOT 'state')
               → fallback: t.get('state', 'Unknown')
  anomalies    → t.get('anomalies', [])
  priority     → t.get('priority', 0)
  status       → NOT reliably stored; compute from anomalies

DLL dict (from dll_parser):
  name         → d.get('name')
  path         → d.get('path')
  base_address → d.get('base_address', 0)    ← exact key (NOT 'base')
  size         → d.get('size', 0)            raw bytes int
  is_signed    → d.get('is_signed')          ← exact key (NOT 'signed')
  entropy      → d.get('entropy', 0.0)
  is_suspicious→ d.get('is_suspicious', False)
  anomalies    → d.get('anomalies', [])
  pid          → d.get('pid', 0)

HANDLE dict (from handle_parser):
  pid          → h.get('pid')
  type         → h.get('type')
  name         → h.get('name','')
  object_address→ h.get('object_address')
  status       → NOT stored; compute from type+name heuristics
  anomalies    → NOT stored; compute from type+name heuristics

BUGS FIXED IN v7
  1. _thread_state(): added 'state_name' as primary key
     (thread_parser stores state_name, not state)
  2. _handle_status(): was always returning "clean" because handles
     have no stored status/is_suspicious.  Now mirrors the tool's
     display-time heuristics (lines 908–932 of extraction tool).
  3. _handle_anomaly(): handles have no stored anomalies.  Compute
     them live from the same heuristics as the extraction tool.
  4. _dll_base(): 'base_address' is now first in lookup order (was 4th).
  5. _dll_signed(): 'is_signed' is now first in lookup order (was 2nd).
  6. _proc_name(): 'image_name' promoted to first in lookup order.
"""

import csv
import os
import re
from datetime import datetime
from collections import defaultdict

from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QFrame, QPushButton, QTableWidget, QTableWidgetItem,
    QLineEdit, QHeaderView, QSizePolicy, QAbstractItemView,
    QStyledItemDelegate, QStyle, QDialog,
)
from PyQt5.QtCore import Qt, QTimer, QRect, QSize
from PyQt5.QtGui  import QFont, QColor, QBrush, QPainter, QPen

from ui.widgets.header_bar   import HeaderBar
from ui.bridge.session       import AppState
from ui.bridge.config_bridge import Config

def _project_root_for_outputs() -> str:
    """Return the current RAMDetect project root for portable page exports.

    This prevents a copied workstation client from writing artifact exports to an
    old absolute path stored in config.yaml, such as another PC's GITHUB folder.
    """
    here = os.path.abspath(__file__)
    cur = os.path.dirname(here)
    while True:
        if os.path.isfile(os.path.join(cur, "main.py")) and os.path.isdir(os.path.join(cur, "ramdetect")):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            # Fallback: artifacts_viewer.py -> pages -> ui -> ramdetect -> project
            return os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
        cur = parent


def _project_reports_dir() -> str:
    reports = os.path.join(_project_root_for_outputs(), "reports")
    os.makedirs(reports, exist_ok=True)
    return reports


try:
    from ramdetect.utils.process_confidence import confidence_label as _j10_conf_label, confidence_tooltip as _j10_conf_tooltip
except Exception:
    def _j10_conf_label(proc):
        return str(proc.get("process_confidence_label") or proc.get("process_confidence") or "Not assessed") if isinstance(proc, dict) else "Not assessed"
    def _j10_conf_tooltip(proc):
        return "Process confidence was not assessed for this session."


TABS = ["Processes", "Threads", "DLLs", "Handles", "PE Sections"]

_TAB_TIPS = {
    "Processes": "Processes are running programs recovered from the memory dump. Use this to inspect PID, parent PID, path, and process-level anomalies.",
    "Threads": "Threads are execution units inside a process. Suspicious thread start addresses can indicate injection or abnormal execution.",
    "DLLs": "DLLs are loaded modules/libraries. Unusual DLL paths, unsigned modules, and LDR-list mismatches can support forensic interpretation.",
    "Handles": "Handles show open references to files, registry keys, events, processes, and other Windows objects.",
    "PE Sections": "PE Sections and mapped regions describe executable memory areas, permissions, entropy, and possible injected code regions. Injection-candidate rows may show N/A for PE-only fields when no PE section bytes were parsed.",
}

_HEADER_TIPS = {
    "PID": "Process Identifier. A unique number used by Windows to identify a process.",
    "PPID": "Parent Process Identifier. Shows which process created or owns this process.",
    "TID": "Thread Identifier. A unique number used to identify a thread inside a process.",
    "ImageFileName": "Raw short process image name recovered from the EPROCESS structure.",
    "Process Name": "Display-friendly process name used by RAMDetect.",
    "Process": "The process associated with this artifact.",
    "Path": "Recovered executable or artifact path. Hover cells to view the full path.",
    "Threads": "Number of threads linked to the process.",
    "Handles": "Number of handles linked to the process.",
    "Status": "Clean means no display-time artifact anomaly. Suspicious means parser or heuristic indicators were found.",
    "Anomaly": "Short explanation of the indicator that made this artifact suspicious.",
    "Start Address": "Memory address where a thread begins execution. Unusual addresses can support injection analysis.",
    "State": "Recovered thread execution/wait state.",
    "DLL Name": "Name of the loaded DLL or module.",
    "Base Address": "Memory base address where the DLL/module was loaded.",
    "Size": "Recovered module size.",
    "Signed": "Whether the DLL appears signed. Unsigned modules may require further review.",
    "Entropy": "Approximate randomness of the data. High entropy may indicate packing, encryption, or obfuscation.",
    "in_load": "Whether the DLL appears in the PEB InLoadOrder list.",
    "in_mem": "Whether the DLL appears in the PEB InMemoryOrder list.",
    "in_init": "Whether the DLL appears in the PEB InInitializationOrder list.",
    "not_in_load": "Yes means the module was not found in the expected InLoadOrder list.",
    "not_in_mem": "Yes means the module was not found in the expected InMemoryOrder list.",
    "not_in_init": "Yes means the module was not found in the expected InInitializationOrder list.",
    "ldr_lists_complete": "Whether the common PEB loader lists appear consistent for the DLL.",
    "ldr_source": "Source or method used to recover the DLL loader-list information.",
    "Type": "Windows object type, such as File, Key, Event, Thread, Process, or Directory.",
    "Handle Name": "Name/path of the referenced Windows object. Hover to view the full value.",
    "Section": "PE section name or mapped memory region label.",
    "Virtual Addr": "Virtual address of the PE section, VAD region, thread start, or injection candidate. N/A means the address was not recoverable from reliable parser evidence.",
    "Virtual Size": "Virtual memory size of the section, VAD region, or matched injection evidence. N/A means no reliable size was observed.",
    "Raw Size": "Raw PE section size when actual PE section bytes were parsed. Injection candidates usually show N/A because they are evidence rows rather than carved PE sections.",
    "Permissions": "Memory protection flags. Writable + executable regions are more suspicious. UNKNOWN means protection was not observed; UNKNOWN(number) means a numeric protection code was recovered but not decoded yet.",
}

_ANOMALY_EXPLAINS = {
    "unexpected path": "Definition: the artifact path does not match the expected normal location for that process, module, or memory artifact.",
    "suspicious path": "Definition: the artifact is located in a path commonly associated with suspicious activity, such as temporary folders, user-writable folders, or unusual system paths.",
    "path anomaly": "Definition: the recovered file path is abnormal, inconsistent, missing, malformed, or not typical for the related artifact.",
    "orphan process": "Definition: a process whose parent process ID cannot be matched to a valid recovered parent process.",
    "parent missing": "Definition: the parent process referenced by the artifact was not recovered or cannot be resolved from the process list.",
    "hidden process": "Definition: a process that appears in one memory structure but is missing or inconsistent in another expected process structure.",
    "code injection": "Definition: code or executable memory appears to have been inserted into a process instead of being loaded through normal execution flow.",
    "reflective dll": "Definition: a DLL-like payload appears to be loaded directly from memory rather than through the normal Windows loader.",
    "hollowing": "Definition: a process may have been created in a legitimate form and then had its original memory replaced with different executable content.",
    "apc injection": "Definition: code execution may be triggered through an Asynchronous Procedure Call queued to a thread.",
    "dropper": "Definition: a program or artifact that acts as an initial-stage component used to place, unpack, or launch another payload.",
    "staged": "Definition: malware behavior divided into multiple phases, where one component prepares, downloads, loads, or triggers another component.",
    "high entropy": "Definition: the data appears highly random, which may indicate packing, encryption, compression, or obfuscation.",
    "unsigned": "Definition: the file or module does not have a verified digital signature.",
    "ldr": "Definition: the module has inconsistency in Windows loader list metadata, such as appearing in one loader list but not another.",
    "not in load": "Definition: the module was not found in the expected PEB InLoadOrder loader list.",
    "not in mem": "Definition: the module was not found in the expected PEB InMemoryOrder loader list.",
    "not in init": "Definition: the module was not found in the expected PEB InInitializationOrder loader list.",
    "rwx": "Definition: the memory region has read, write, and execute permissions at the same time.",
    "persistence mechanism": "Definition: the artifact references a location commonly used to make a program start automatically after reboot or user login.",
    "service manipulation": "Definition: the artifact references Windows service configuration areas that can be used to install, change, or persist services.",
    "security database access": "Definition: the artifact references sensitive Windows security database files or registry areas such as SAM or SECURITY.",
    "sensitive file access": "Definition: the artifact references files that may contain passwords, credentials, keys, or other sensitive data.",
    "suspicious temp execution": "Definition: an executable or script-like object is associated with a temporary directory.",
    "suspicious temp": "Definition: an artifact is associated with a temporary directory, which is often used for unpacking or staging files.",
    "service": "Definition: the artifact is related to a Windows service or service configuration object.",
    "credential": "Definition: the artifact references credential-related data such as passwords, keys, tokens, or authentication files.",
    "shellcode": "Definition: small executable code designed to run directly from memory, often used after exploitation or injection.",
}

def _anomaly_tooltip(anomaly_text: str) -> str:
    """Return definitions of displayed anomaly labels.

    This tooltip explains the meaning of the anomaly term only. It does not
    change or imply the ML verdict.
    """
    raw = str(anomaly_text or "").strip()
    if not raw or raw == "-":
        return "Definition: no anomaly label was produced for this artifact."

    low = raw.lower()
    matches = []
    used = set()
    for key, definition in _ANOMALY_EXPLAINS.items():
        if key in low and key not in used:
            matches.append(definition)
            used.add(key)

    if not matches:
        matches.append(
            "Definition: this is a RAMDetect artifact-level anomaly label. "
            "It describes an unusual property of the artifact for analyst review."
        )

    return "Anomaly: " + raw + "\n\n" + "\n".join(f"• {m}" for m in matches[:3])

_THREAD_STATE_EXPLAINS = {
    "initialized": "Initialized: the thread object was created but has not started running.",
    "ready": "Ready: the thread is ready to run and waiting for CPU scheduling.",
    "running": "Running: the thread is currently executing.",
    "standby": "Standby: the thread has been selected to run next.",
    "terminated": "Terminated: the thread has finished execution.",
    "waiting": "Waiting: the thread is waiting for an event, I/O, timer, lock, or another condition.",
    "transition": "Transition: the thread is waiting for memory/resources before it can run.",
    "deferredready": "Deferred Ready: the thread is ready but scheduling is temporarily deferred.",
    "gatewaitobsol": "Gate Wait: older/obsolete Windows wait state naming.",
}

_PERMISSION_EXPLAINS = {
    "NOACCESS": "NOACCESS: this memory region cannot be read, written, or executed.",
    "READONLY": "READONLY: readable memory only; usually lower risk.",
    "READWRITE": "READWRITE: readable and writable memory. Normal for data, but not executable.",
    "WRITECOPY": "WRITECOPY: copy-on-write memory. Common for mapped images/sections.",
    "EXECUTE": "EXECUTE: executable memory only.",
    "EXECUTE_READ": "EXECUTE_READ: executable and readable memory. Common for normal code sections.",
    "EXECUTE_READWRITE": "EXECUTE_READWRITE / RWX: readable, writable, and executable. This should be reviewed because it can support injected code or shellcode.",
    "EXECUTE_WRITECOPY": "EXECUTE_WRITECOPY: executable copy-on-write memory; review with context.",
}

def _thread_state_tooltip(state_text: str) -> str:
    """Explain thread state labels/codes for beginner users."""
    raw = str(state_text or "").strip()
    if not raw:
        return "Thread state was not recovered."
    low = raw.lower()
    for key, desc in _THREAD_STATE_EXPLAINS.items():
        if key in low:
            return raw + "\n\n" + desc
    try:
        m = re.search(r"\((\d+)\)", raw)
    except NameError:
        m = None
    if "unknown" in low and m:
        return (
            f"{raw}\n\nRAMDetect recovered numeric thread state code {m.group(1)}, "
            "but it is not mapped to a named Windows thread state in the current parser. "
            "This is not automatically suspicious by itself; review it together with "
            "the thread start address and anomaly column."
        )
    if "unknown" in low:
        return (
            raw + "\n\nThe parser could not map this thread state to a known label. "
            "Review with the start address and anomaly column."
        )
    return raw + "\n\nThread state recovered from the memory artifact."


_EXCEL_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")

def _is_invalid_xml_char(ch: str) -> bool:
    """Return True when a character cannot be safely written into XLSX XML."""
    cp = ord(ch)

    # XML 1.0 legal character ranges.
    legal_xml_10 = (
        cp in (0x09, 0x0A, 0x0D)
        or 0x20 <= cp <= 0xD7FF
        or 0xE000 <= cp <= 0xFFFD
        or 0x10000 <= cp <= 0x10FFFF
    )
    if not legal_xml_10:
        return True

    # Unicode noncharacters can still break Excel/XML parsers. RAM dumps may
    # contain them when raw bytes are decoded into text.
    if 0xFDD0 <= cp <= 0xFDEF:
        return True
    if (cp & 0xFFFF) in (0xFFFE, 0xFFFF):
        return True

    # Extra guard for the U+FFF0..U+FFFF noncharacter/special range observed
    # in corrupted DLL path artifacts.
    if 0xFFF0 <= cp <= 0xFFFF:
        return True

    return False

def _strip_export_unsafe_chars(value) -> str:
    """Remove characters unsafe for XML-based Excel output."""
    if value is None:
        return ""
    s = str(value)
    return "".join(ch for ch in s if not _is_invalid_xml_char(ch))

def _excel_safe_cell(value):
    """Return a value safe for openpyxl worksheet cells.

    Memory artifacts can contain raw bytes decoded as XML-invalid Unicode
    characters. Excel/OpenXML cannot store those characters, so they must be
    removed before writing the workbook.
    """
    if value is None:
        return ""

    # Keep basic numeric values as-is.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value

    s = _strip_export_unsafe_chars(value)

    # Avoid Excel interpreting forensic strings as formulas.
    if s.startswith(("=", "+", "-", "@")):
        s = "'" + s

    return s

def _csv_safe_cell(value):
    """Return text safe for CSV without adding Excel-only formula prefixes."""
    return _strip_export_unsafe_chars(value)

def _permission_tooltip(perms: str) -> str:
    """Explain PE/VAD permission values without implying missing values are parser failures."""
    raw = str(perms or "").strip()
    if not raw or raw in {"?", "N/A"}:
        return (
            "Memory permission value was not available for this row. If this is an "
            "injection candidate, it may have been promoted from thread/path/VAD "
            "evidence rather than from a parsed PE section. RAMDetect keeps the "
            "evidence visible but avoids guessing missing values."
        )
    up = raw.upper().replace(" ", "_")
    if up in {"UNKNOWN", "UNKNOWN_PROTECTION", "NOT_OBSERVED"}:
        return (
            raw + "\n\n"
            "Protection was not observed from a matching VAD or detector row. "
            "The artifact can still be valid injection evidence, but RAMDetect "
            "does not infer a permission value when the source is unavailable."
        )
    if up in _PERMISSION_EXPLAINS:
        return raw + "\n\n" + _PERMISSION_EXPLAINS[up]

    m = re.search(r"UNKNOWN\((\d+)\)", up)
    if m:
        return (
            raw + "\n\n"
            f"RAMDetect recovered numeric memory protection code {m.group(1)}, "
            "but this code is not currently mapped to a readable permission label. "
            "This does not invalidate the finding; it means the permission decoder "
            "needs a mapping update for this value."
        )

    parts = []
    if "R" in up or "READ" in up:
        parts.append("readable")
    if "W" in up or "WRITE" in up:
        parts.append("writable")
    if "X" in up or "EXECUTE" in up:
        parts.append("executable")
    if parts:
        msg = f"This region appears to be {', '.join(parts)}."
        if ("W" in up or "WRITE" in up) and ("X" in up or "EXECUTE" in up):
            msg += " Writable + executable memory should be reviewed for possible injection."
        return raw + "\n\n" + msg
    return raw + "\n\nMemory protection/permission string recovered by RAMDetect."

# ─────────────────────────────────────────────────────────────────────────────
# Toast notification  (auto-closing pop-out window)
# ─────────────────────────────────────────────────────────────────────────────

class _ToastDialog(QWidget):
    """Small frameless pop-up that auto-closes after `ms` milliseconds."""

    def __init__(self, title: str, msg: str, ms: int = 4000, parent=None):
        super().__init__(parent, Qt.Tool | Qt.FramelessWindowHint |
                         Qt.WindowStaysOnTopHint)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setAttribute(Qt.WA_DeleteOnClose)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        box = QFrame()
        is_err = title.startswith("X") or "failed" in title.lower()
        border = "#f87171" if is_err else "#34d399"
        box.setStyleSheet(
            f"QFrame{{background:#1e293b;border:2px solid {border};"
            f"border-radius:12px;}}"
        )
        bv = QVBoxLayout(box)
        bv.setContentsMargins(20, 14, 20, 14)
        bv.setSpacing(6)

        t = QLabel(title)
        t.setStyleSheet(
            f"color:{border};font-size:13px;font-weight:700;"
            "background:transparent;border:none;"
        )
        bv.addWidget(t)

        m = QLabel(msg)
        m.setStyleSheet(
            "color:#94a3b8;font-size:11px;"
            "background:transparent;border:none;"
        )
        m.setWordWrap(True)
        bv.addWidget(m)

        outer.addWidget(box)
        self.adjustSize()
        QTimer.singleShot(ms, self.close)

    def show_near(self, parent_widget):
        """Position bottom-right of the parent window."""
        if parent_widget and parent_widget.window():
            pw  = parent_widget.window()
            geo = pw.geometry()
            x   = geo.right()  - self.width()  - 24
            y   = geo.bottom() - self.height() - 24
            self.move(x, y)
        self.show()
_HEADER_H = 42
_BADGE_W  = 130
_BADGE_H  = 26


# ─────────────────────────────────────────────────────────────────────────────
# Status badge delegate
# ─────────────────────────────────────────────────────────────────────────────

class _StatusDelegate(QStyledItemDelegate):
    _SUSP_BG  = QColor(220, 38,  38,  210)
    _SUSP_BRD = QColor(239, 68,  68,  255)
    _CLEN_BG  = QColor(5,   150, 105, 210)
    _CLEN_BRD = QColor(16,  185, 129, 255)
    _TEXT_COL = QColor(255, 255, 255, 255)
    _FONT = QFont("Segoe UI", 9)
    _FONT.setBold(True)

    def paint(self, painter, option, index):
        painter.save()
        painter.setRenderHint(QPainter.Antialiasing, True)
        if option.state & QStyle.State_Selected:
            painter.fillRect(option.rect, QColor(6, 182, 212, 40))
        elif option.state & QStyle.State_MouseOver:
            painter.fillRect(option.rect, QColor(51, 65, 85, 70))
        text    = index.data(Qt.DisplayRole) or ""
        is_susp = "SUSPICIOUS" in text
        x    = option.rect.x() + max(0, (option.rect.width()  - _BADGE_W) // 2)
        y    = option.rect.y() +        (option.rect.height() - _BADGE_H) // 2
        pill = QRect(x, y, _BADGE_W, _BADGE_H)
        painter.setBrush(QBrush(self._SUSP_BG  if is_susp else self._CLEN_BG))
        painter.setPen(QPen(self._SUSP_BRD if is_susp else self._CLEN_BRD, 1.5))
        painter.drawRoundedRect(pill, 8, 8)
        painter.setPen(self._TEXT_COL)
        painter.setFont(self._FONT)
        painter.drawText(pill, Qt.AlignCenter, text)
        painter.restore()

    def sizeHint(self, option, index):
        return QSize(_BADGE_W + 20, 44)


# ─────────────────────────────────────────────────────────────────────────────
# Artifact count helpers  (mirror extraction_tool.count_artifacts_by_pid)
# ─────────────────────────────────────────────────────────────────────────────

def _count_by_pid(items: list) -> dict:
    counts: dict = defaultdict(int)
    for item in items:
        pid_raw = item.get("pid") or item.get("PID")
        if pid_raw is None:
            continue
        try:
            counts[int(str(pid_raw))] += 1
        except (ValueError, TypeError):
            pass
    return dict(counts)


# ─────────────────────────────────────────────────────────────────────────────
# Field-name helpers — aligned to exact extraction-tool keys
# ─────────────────────────────────────────────────────────────────────────────

def _pick(d: dict, *keys, default="?") -> str:
    """Return first truthy non-empty value; 0/'0' treated as missing."""
    for k in keys:
        v = d.get(k)
        if v is not None and str(v).strip() not in ("", "None", "0", "0x0"):
            return str(v)
    return default


# ── PROCESS ───────────────────────────────────────────────────────────────────
# Extraction tool uses: image_name, pid, ppid, path, anomalies, is_suspicious

def _proc_image_file_name(p) -> str:
    # Raw 15-byte EPROCESS.ImageFileName field, used for Volatility-parity export.
    return _pick(p, "image_file_name", "ImageFileName", "ImageFileNameRaw",
                 "image_name", "name", "process_name", "proc_name",
                 "exe_name", "filename")

def _proc_name(p) -> str:
    # Display-friendly RAMDetect name.
    return _pick(p, "image_name", "process_name", "name", "ImageFileName",
                 "proc_name", "exe_name", "filename")

def _proc_pid(p) -> str:
    return _pick(p, "pid", "PID", "UniqueProcessId", "process_id")

def _proc_ppid(p) -> str:
    return _pick(p, "ppid", "PPID", "InheritedFromUniqueProcessId",
                 "parent_pid", "parent_id", default="-")

def _proc_path(p) -> str:
    return _pick(p, "path", "exe_path", "ImagePathName", "full_path",
                 "image_path", "executable", default="-")

def _proc_status(p) -> str:
    """
    Mirror determine_status() from extraction tool (line 384–388):
      anomalies present → SUSPICIOUS
      status == 'SUSPICIOUS' → SUSPICIOUS
      is_suspicious == True → SUSPICIOUS
    """
    if (p.get("status", "").upper() == "SUSPICIOUS"
            or p.get("is_suspicious")
            or p.get("anomalies")):
        return "suspicious"
    return "clean"

def _proc_anomaly(p) -> str:
    anoms = p.get("anomalies", [])
    if not anoms:
        return "-"
    parts = [
        a.get("type", str(a)) if isinstance(a, dict)
        else str(a).split(":")[0].strip()
        for a in anoms[:2]
    ]
    return " + ".join(parts)


# ── THREAD ────────────────────────────────────────────────────────────────────
# Extraction tool uses: tid, pid, start_address, state_name/state, anomalies

def _thread_tid(t) -> str:
    return _pick(t, "tid", "TID", "ThreadId", "UniqueThread", "thread_id")

def _thread_pid(t) -> str:
    return _pick(t, "pid", "PID", "owning_pid", "process_id")

def _thread_process(t, pid_map=None) -> str:
    """Process name not stored in thread dict — resolve via pid_map."""
    name = _pick(t, "process", "process_name", "proc_name",
                 "image_name", "ImageFileName",
                 "owner", "owner_name", default=None)
    if name and name != "?":
        return name
    if pid_map:
        pid_raw = t.get("pid") or t.get("PID")
        if pid_raw is not None:
            try:
                return pid_map.get(int(str(pid_raw))) or "?"
            except (ValueError, TypeError):
                pass
    return "?"

def _thread_start_addr(t) -> str:
    """
    Extraction tool key is 'start_address' (line 617: t.get('start_address', 0)).
    Format as 0x hex.
    """
    raw = _pick(t, "start_address", "start_addr",
                "StartAddress", "win32_start_addr", "Win32StartAddress",
                "ethread_start_addr", "thread_start", "start", default=None)
    if not raw:
        return "?"
    try:
        n = int(raw, 0) if isinstance(raw, str) and raw.startswith("0x") else int(raw)
        return f"0x{n:016X}"
    except (ValueError, TypeError):
        return str(raw)

def _thread_state(t) -> str:
    """
    Extraction tool primary key is 'state_name' (line 618:
      t.get('state_name', t.get('state', 'Unknown'))).
    Must try state_name FIRST.
    """
    return _pick(t, "state_name", "state", "State", "thread_state", "wait_reason")

def _thread_status(t) -> str:
    """
    Thread status is NOT reliably stored — mirror determine_status().
    """
    if (t.get("status", "").upper() == "SUSPICIOUS"
            or t.get("is_suspicious")
            or t.get("anomalies")):
        return "suspicious"
    return "clean"

def _thread_anomaly(t) -> str:
    anoms = t.get("anomalies", [])
    if not anoms:
        return t.get("anomaly", "-") or "-"
    if isinstance(anoms, list):
        parts = [
            a.get("type", str(a)) if isinstance(a, dict)
            else str(a).split(":")[0].strip()
            for a in anoms[:2]
        ]
        return " + ".join(parts)
    return str(anoms)


# ── DLL ───────────────────────────────────────────────────────────────────────
# Extraction tool uses: pid, process, name, path, base_address, size, is_signed, entropy

def _dll_pid(d) -> str:
    return _pick(d, "pid", "PID", default="-")

def _dll_process(d) -> str:
    return _pick(d, "process", "Process", "process_name", "ImageFileName", default="-")

def _dll_name(d) -> str:
    return _pick(d, "name", "dll_name", "BaseDllName", "FullDllName",
                 "module_name", "filename")

def _dll_path(d) -> str:
    return _pick(d, "path", "full_path", "FullDllName", "image_path",
                 "dll_path", default="-")

def _dll_base(d) -> str:
    """Primary key is 'base_address' (line 1122: d.get('base_address', 0))."""
    raw = _pick(d, "base_address", "base", "base_addr",
                "DllBase", "load_address", default=None)
    if not raw:
        return "?"
    try:
        n = int(raw, 0) if isinstance(raw, str) and raw.startswith("0x") else int(raw)
        return f"0x{n:X}"
    except (ValueError, TypeError):
        return str(raw)

def _dll_size(d) -> str:
    """Stored as raw bytes integer (line 1122: d.get('size', 0))."""
    raw = _pick(d, "size", "SizeOfImage", "dll_size", "image_size", default=None)
    if not raw:
        return "?"
    try:
        n = int(raw, 0) if isinstance(raw, str) and raw.startswith("0x") else int(raw)
        return f"{n/1048576:.1f} MB" if n >= 1048576 else f"{n//1024} KB"
    except (ValueError, TypeError):
        return str(raw)

def _dll_signed(d) -> bool:
    """Primary key is 'is_signed' (line 1123: d.get('is_signed'))."""
    v = d.get("is_signed")
    if v is None:
        v = d.get("signed")
    if v is None:
        return True   # unknown → assume signed
    return v if isinstance(v, bool) else str(v).lower() in ("true","yes","1")

def _dll_entropy(d) -> float:
    try:
        return float(d.get("entropy", d.get("dll_entropy", 0.0)))
    except:
        return 0.0

def _dll_status(d) -> str:
    """Mirror extraction tool line 1119: is_suspicious OR anomalies → SUSPICIOUS."""
    if (d.get("is_suspicious", False) or d.get("anomalies")
            or d.get("status", "").upper() == "SUSPICIOUS"):
        return "suspicious"
    return "clean"


def _dll_bool_field(d, *keys) -> str:
    """Render LDR membership booleans for table/CSV export.

    Returns Yes/No for known booleans and '-' when the field was not supplied
    by the parser. This prevents old DLL rows from being displayed as false
    mismatches.
    """
    for k in keys:
        if k in d and d.get(k) is not None:
            v = d.get(k)
            if isinstance(v, bool):
                return "Yes" if v else "No"
            s = str(v).strip().lower()
            if s in ("true", "1", "yes", "y"):
                return "Yes"
            if s in ("false", "0", "no", "n"):
                return "No"
            if s not in ("", "none", "null", "?"):
                return str(v)
    return "-"

def _dll_text_field(d, *keys, default="-") -> str:
    return _pick(d, *keys, default=default)


# ── HANDLE ────────────────────────────────────────────────────────────────────
# Raw handle dict only stores: pid, type, name, object_address
# status and anomalies are computed at display time (lines 908–932 of tool)

def _handle_type(h) -> str:
    return _pick(h, "type", "handle_type", "Type", "object_type")

def _handle_name(h) -> str:
    return _pick(h, "name", "handle_name", "Name", "object_name", "path")

def _handle_pid(h) -> str:
    return _pick(h, "pid", "PID", "process_id")

def _handle_process(h, pid_map=None) -> str:
    """Process name not stored — resolve via pid_map."""
    name = _pick(h, "process", "process_name", "image_name",
                 "ImageFileName", "proc_name", default=None)
    if name and name != "?":
        return name
    if pid_map:
        pid_raw = h.get("pid") or h.get("PID")
        if pid_raw is not None:
            try:
                return pid_map.get(int(str(pid_raw))) or "?"
            except (ValueError, TypeError):
                pass
    return "?"

def _handle_compute(h: dict) -> tuple:
    """
    Compute (status, anomaly_str) for a handle using the same heuristics
    as show_detailed_handles_table() lines 908–932 of the extraction tool.

    Raw handle dicts NEVER have a stored status/anomalies field — the tool
    derives them live from type+name at display time.
    """
    h_type   = h.get("type", "")
    h_name   = h.get("name", "")
    name_low = h_name.lower()
    anomalies = []

    # --- Registry checks (tool lines 913–918) ---
    if h_type == "Key":
        if "currentversion\\run" in name_low or "runonce" in name_low:
            anomalies.append("Persistence Mechanism")
        if "services" in name_low and "currentcontrolset" in name_low:
            anomalies.append("Service Manipulation")

    # --- File checks (tool lines 920–926) ---
    if h_type == "File":
        if "sam" in name_low or ("security" in name_low and "config" in name_low):
            anomalies.append("Security Database Access")
        if ("password" in name_low or "credential" in name_low
                or ".kdbx" in name_low):
            anomalies.append("Sensitive File Access")
        if "temp" in name_low and (".exe" in name_low or ".bat" in name_low):
            anomalies.append("Suspicious Temp Execution")

    # If stored status / anomalies fields exist, honour them too
    if h.get("status", "").upper() == "SUSPICIOUS" or h.get("is_suspicious"):
        if not anomalies:
            anomalies.append("Flagged by parser")

    stored_anoms = h.get("anomalies", [])
    if stored_anoms:
        if isinstance(stored_anoms, list):
            anomalies.extend(str(a) for a in stored_anoms)
        else:
            anomalies.append(str(stored_anoms))

    status    = "suspicious" if anomalies else "clean"
    anom_str  = "; ".join(anomalies[:2]) if anomalies else "-"
    return status, anom_str


# ── PE SECTION ────────────────────────────────────────────────────────────────


def _pe_proc(s, pid_map=None) -> str:
    """Resolve process name using pid_map, fallback to direct keys."""
    name = _pick(s, "process", "process_name", "image_name",
                 "ImageFileName", "proc_name", default=None)
    if name and name != "?":
        return name
    if pid_map:
        pid_raw = s.get("pid") or s.get("PID")
        if pid_raw is not None:
            try:
                return pid_map.get(int(str(pid_raw))) or "?"
            except (ValueError, TypeError):
                pass
    return "?"

def _pe_section(s) -> str:
    # ADDED "type" so VAD regions show as 'Mapped' or 'Private'
    return _pick(s, "section", "section_name", "Name", "name", "type")

def _pe_missing(value) -> bool:
    if value is None:
        return True
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) == 0
    text = str(value).strip()
    return text == "" or text.lower() in {"none", "n/a", "na", "?", "-", "0", "0x0"}

def _pe_first_value(s: dict, *keys):
    for key in keys:
        value = s.get(key) if isinstance(s, dict) else None
        if not _pe_missing(value):
            return value
    return None

def _pe_hex_value(value) -> str:
    if _pe_missing(value):
        return "?"
    try:
        if isinstance(value, str):
            text = value.strip()
            n = int(text, 16) if text.lower().startswith("0x") else int(float(text))
        else:
            n = int(value)
        return f"0x{n:08X}"
    except (ValueError, TypeError):
        return str(value)

def _pe_hex(s, *keys) -> str:
    return _pe_hex_value(_pe_first_value(s, *keys))

def _pe_na_reason(s: dict, field_name: str) -> str:
    source_keys = {
        "Virtual Addr": ("address_source", "start_address_source", "vad_source"),
        "Virtual Size": ("size_source", "region_size_source", "vad_source"),
        "Raw Size": ("raw_size_source", "pe_source"),
        "Entropy": ("entropy_source", "pe_source"),
    }.get(field_name, ())
    source = _pe_first_value(s, *source_keys) if source_keys else None
    if source:
        return f" Evidence source: {source}."
    return ""

def _pe_is_injection_candidate(s: dict) -> bool:
    """Return True for promoted injection/malfind rows shown in the PE tab.

    J9B: RAMDetect can display InjectionDetector rows in the PE Sections tab
    even when their type/section is not literally named "injection candidate".
    Examples include "Suspicious Thread Location" and generic "Code Injection"
    anomaly rows.  Treat these as injection evidence so missing PE-only fields
    render as N/A with explanatory tooltips instead of a bare question mark.
    """
    if not isinstance(s, dict):
        return False
    if s.get("i2f_feature_flag") or s.get("source") == "ramdetect_i2f_kernel_start_plus_suspicious_path":
        return True
    if s.get("is_suspicious") and any(k in s for k in (
        "address", "virtual_address", "start_address", "thread_start_address",
        "vad_start", "protection", "vad_protection", "matched_vad_protection",
        "protection_source", "address_source", "size_source",
    )):
        return True
    text = " ".join(str(s.get(k, "")) for k in (
        "type", "name", "section", "section_name", "anomaly", "anomalies",
        "source", "policy", "details", "status_reason", "reason",
    )).lower()
    injection_terms = (
        "injection candidate", "code injection", "malfind", "suspicious thread location",
        "kernel start", "suspicious path injection", "thread start in private",
        "private executable", "private rwx", "unbacked", "shellcode", "hollow",
        "mz/pe in private", "pe in anonymous",
    )
    return any(term in text for term in injection_terms)

def _pe_na_tooltip(field_name: str, row: dict | None = None) -> str:
    extra = _pe_na_reason(row or {}, field_name) if row else ""
    if field_name == "Raw Size":
        return (
            "Raw Size is not available because this item is not a parsed PE section "
            "with carved raw section bytes. This is normal for injection evidence rows."
            + extra
        )
    if field_name == "Entropy":
        return (
            "Entropy requires actual section bytes. This row did not contain reliable "
            "bytes for entropy calculation, so RAMDetect shows N/A instead of guessing."
            + extra
        )
    return (
        f"{field_name} is not available for this row because RAMDetect did not "
        "recover a reliable value from a matching VAD, thread, or detector row. "
        "The evidence is preserved, but missing fields are not inferred."
        + extra
    )

def _pe_hex_display(s: dict, field_name: str, *keys) -> str:
    value = _pe_hex(s, *keys)
    # J9B: suspicious injection evidence rows are not normal parsed PE sections.
    # When PE-only fields are absent, show N/A rather than a bare question mark.
    if value == "?" and (_pe_is_injection_candidate(s) or s.get("is_suspicious")):
        return "N/A"
    return value

def _pe_virtual_addr_display(s: dict) -> str:
    return _pe_hex_display(
        s, "Virtual Addr",
        "address", "virtual_address", "start_address", "thread_start_address",
        "vad_start", "start", "va", "virt_addr", "VirtualAddress", "base_address",
    )

def _pe_virtual_size_display(s: dict) -> str:
    return _pe_hex_display(
        s, "Virtual Size",
        "region_size", "vad_size", "size", "virt_size", "VirtualSize",
    )

def _pe_raw_size_display(s: dict) -> str:
    # Raw size is a PE-section-only value. Do not use VAD end address as raw size.
    return _pe_hex_display(s, "Raw Size", "raw_size", "SizeOfRawData")

def _pe_entropy(s) -> float:
    try:
        return float(s.get("entropy", s.get("section_entropy", 0.0)))
    except Exception:
        return 0.0

def _pe_entropy_display(s: dict) -> str:
    has_entropy = any(k in s and str(s.get(k)).strip() not in {"", "None", "?", "N/A"}
                      for k in ("entropy", "section_entropy"))
    if _pe_is_injection_candidate(s) and not has_entropy:
        return "N/A"
    return f"{_pe_entropy(s):.1f}"

def _pe_perms(s) -> str:
    value = _pe_first_value(
        s,
        "protection", "vad_protection", "matched_vad_protection",
        "page_protection", "memory_protection", "permissions", "permission",
        "perms", "Characteristics",
    )
    if value is None and _pe_is_injection_candidate(s):
        return "UNKNOWN"
    return str(value) if value is not None else "?"

def _pe_status(s) -> str:
    if (s.get("status", "").upper() == "SUSPICIOUS"
            or s.get("is_suspicious") or s.get("anomalies")):
        return "suspicious"
    return "clean"


# ── Process tree flattening ────────────────────────────────────────────────────

def _flat_procs(procs, depth=0):
    out = []
    for p in procs:
        out.append((p, depth))
        out.extend(_flat_procs(p.get("children", []), depth + 1))
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Sorting
# ─────────────────────────────────────────────────────────────────────────────

def _sort_key(value: str):
    v = str(value).strip()
    if v.startswith(("0x","0X")):
        try: return (0, int(v, 16))
        except ValueError: pass
    try: return (0, int(v))
    except ValueError: pass
    try: return (0, float(v))
    except ValueError: pass
    return (1, v.lower())

def _sorted_rows(rows, getters, sort_col, ascending):
    if sort_col < 0 or sort_col >= len(getters): return rows
    getter = getters[sort_col]
    if getter is None: return rows
    return sorted(rows, key=lambda r: _sort_key(getter(r)), reverse=not ascending)


# ─────────────────────────────────────────────────────────────────────────────
# UI helpers
# ─────────────────────────────────────────────────────────────────────────────

def _card() -> QFrame:
    f = QFrame()
    f.setStyleSheet(
        "QFrame{background:rgba(15,23,42,0.55);"
        "border:1px solid rgba(51,65,85,0.55);border-radius:14px;}"
    )
    return f

def _lbl(text, size=12, bold=False, color="#94a3b8") -> QLabel:
    l = QLabel(text)
    l.setStyleSheet(
        f"color:{color};background:transparent;border:none;"
        f"font-size:{size}px;font-weight:{'700' if bold else '400'};"
    )
    l.setAlignment(Qt.AlignCenter)
    return l

def _compact_tooltip(text: str, limit: int = 600) -> str:
    """Keep tooltips useful without making Qt render very large popups."""
    s = str(text or "")
    if len(s) <= limit:
        return s
    return s[:limit] + "\n... [truncated for UI performance]"


# ─────────────────────────────────────────────────────────────────────────────
# Page
# ─────────────────────────────────────────────────────────────────────────────

class ArtifactsViewerPage(QWidget):

    def __init__(self, main_window, parent=None):
        super().__init__(parent)
        self.mw          = main_window
        self._active_tab = 0
        self._susp_only  = False
        self._sort_col   = -1
        self._sort_asc   = True
        self.setStyleSheet("background:transparent;")

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        hdr = HeaderBar(show_action=False, show_dashboard=True, show_results=True)
        hdr.dashboard_clicked.connect(lambda: self.mw.navigate_to("dashboard"))
        hdr.settings_clicked.connect(lambda: self.mw.navigate_to("settings"))
        hdr.results_clicked.connect(lambda: self.mw.navigate_to("results"))
        root.addWidget(hdr)

        body = QVBoxLayout()
        body.setContentsMargins(32, 20, 32, 20)
        body.setSpacing(14)
        body.addWidget(self._build_tab_row())
        body.addWidget(self._build_table_card(), stretch=1)
        root.addLayout(body, stretch=1)

    # ─────────────────────────────────────────────────────────────────────────
    # Tab row
    # ─────────────────────────────────────────────────────────────────────────

    def _build_tab_row(self) -> QWidget:
        w = QWidget()
        w.setStyleSheet("background:transparent;")
        row = QHBoxLayout(w)
        row.setContentsMargins(0, 0, 0, 0)
        row.setSpacing(6)

        self._tab_btns = []
        for i, name in enumerate(TABS):
            btn = QPushButton(name)
            btn.setFixedHeight(38)
            btn.setCursor(Qt.PointingHandCursor)
            btn.setStyleSheet(self._tab_style(i == 0))
            btn.setToolTip(_TAB_TIPS.get(name, "View this artifact category."))
            btn.clicked.connect(lambda _, idx=i: self._switch_tab(idx))
            self._tab_btns.append(btn)
            row.addWidget(btn)

        row.addStretch()

        self._susp_btn = QPushButton("⚠  Suspicious")
        self._susp_btn.setCheckable(True)
        self._susp_btn.setFixedHeight(36)
        self._susp_btn.setCursor(Qt.PointingHandCursor)
        self._susp_btn.setToolTip("Show only artifacts with suspicious indicators.")
        self._susp_btn.setStyleSheet(self._susp_style(False))
        self._susp_btn.clicked.connect(self._toggle_susp)
        row.addWidget(self._susp_btn)

        search_frame = QFrame()
        search_frame.setFixedHeight(36)
        search_frame.setStyleSheet(
            "QFrame{background:rgba(30,41,59,0.6);"
            "border:1px solid rgba(51,65,85,0.7);border-radius:9px;}"
        )
        sf = QHBoxLayout(search_frame)
        sf.setContentsMargins(10, 0, 10, 0)
        sf.setSpacing(6)
        srch_ic = QLabel("🔍")
        srch_ic.setStyleSheet("color:#64748b;background:transparent;border:none;font-size:12px;")
        sf.addWidget(srch_ic)
        self.search = QLineEdit()
        self.search.setPlaceholderText("Search artifacts...")
        self.search.setToolTip("Search visible artifacts by PID, process name, path, status, or anomaly.")
        self.search.setFixedWidth(190)
        self.search.setStyleSheet(
            "QLineEdit{background:transparent;border:none;color:#e2e8f0;font-size:12px;}"
            "QLineEdit::placeholder{color:#475569;}"
        )
        self.search.textChanged.connect(self._apply_row_filter)
        sf.addWidget(self.search)
        row.addWidget(search_frame)

        export_btn = QPushButton("⬇ Export Tab")
        export_btn.setFixedHeight(36)
        export_btn.setCursor(Qt.PointingHandCursor)
        export_btn.setToolTip("Choose CSV or themed XLSX export for the current artifact tab.")
        export_btn.setStyleSheet(
            "QPushButton{background:rgba(30,41,59,0.6);"
            "border:1px solid rgba(51,65,85,0.6);border-radius:9px;"
            "color:#cbd5e1;padding:0 14px;font-size:12px;font-weight:700;}"
            "QPushButton:hover{background:rgba(6,182,212,0.12);color:#ffffff;"
            "border:1px solid rgba(6,182,212,0.35);}"
            "QPushButton:pressed{background:rgba(15,23,42,0.8);}"
        )
        export_btn.clicked.connect(self._show_export_choice_dialog)
        row.addWidget(export_btn)

        xlsx_btn = QPushButton("📊 All XLSX")
        xlsx_btn.setFixedHeight(36)
        xlsx_btn.setCursor(Qt.PointingHandCursor)
        xlsx_btn.setToolTip("Export all artifact tabs into one Excel workbook with separate sheets.")
        xlsx_btn.setStyleSheet(
            "QPushButton{background:rgba(6,182,212,0.10);"
            "border:1px solid rgba(6,182,212,0.35);border-radius:9px;"
            "color:#67e8f9;padding:0 12px;font-size:12px;font-weight:700;}"
            "QPushButton:hover{background:rgba(6,182,212,0.18);color:#ffffff;}"
            "QPushButton:pressed{background:rgba(8,145,178,0.24);}"
        )
        xlsx_btn.clicked.connect(self._export_excel_all)
        row.addWidget(xlsx_btn)

        self._exp_lbl = QLabel("")
        self._exp_lbl.setStyleSheet(
            "color:#34d399;background:transparent;border:none;font-size:11px;"
        )
        row.addWidget(self._exp_lbl)

        self._refresh_tab_styles()
        return w

    @staticmethod
    def _tab_style(active: bool) -> str:
        if active:
            return ("QPushButton{background:#06b6d4;color:#ffffff;"
                    "border-radius:9px;padding:4px 16px;"
                    "font-size:12px;font-weight:700;border:none;}")
        return ("QPushButton{background:rgba(30,41,59,0.5);color:#94a3b8;"
                "border:1px solid rgba(51,65,85,0.5);border-radius:9px;"
                "padding:4px 14px;font-size:12px;}"
                "QPushButton:hover{background:rgba(51,65,85,0.6);color:#e2e8f0;}")

    @staticmethod
    def _susp_style(active: bool) -> str:
        if active:
            return ("QPushButton{background:rgba(248,113,113,0.18);color:#f87171;"
                    "border:1px solid rgba(248,113,113,0.45);border-radius:9px;"
                    "padding:0 14px;font-size:12px;font-weight:600;}")
        return ("QPushButton{background:rgba(30,41,59,0.5);color:#94a3b8;"
                "border:1px solid rgba(51,65,85,0.5);border-radius:9px;"
                "padding:0 14px;font-size:12px;}"
                "QPushButton:hover{background:rgba(248,113,113,0.10);color:#f87171;"
                "border-color:rgba(248,113,113,0.35);}")

    def _refresh_tab_styles(self):
        sess = AppState.current
        
        # FIX: Combine pe_sections and injections for the total count!
        total_sections = list(sess.pe_sections)
        if hasattr(sess, "injections") and sess.injections:
            total_sections.extend(sess.injections)

        artifact_lists = [
            [p for p,_ in _flat_procs(sess.processes)],
            sess.threads, sess.dlls, sess.handles, total_sections, # Updated here!
        ]
        
        for i, btn in enumerate(self._tab_btns):
            btn.setText(f"{TABS[i]}  ({len(artifact_lists[i])})")
            btn.setStyleSheet(self._tab_style(i == self._active_tab))

    def _switch_tab(self, idx: int):
        self._active_tab = idx
        self._sort_col   = -1
        self._sort_asc   = True
        self.search.clear()
        self._exp_lbl.setText("")
        self._refresh_tab_styles()
        self._populate_table()

    def _toggle_susp(self, checked: bool):
        self._susp_only = checked
        self._susp_btn.setStyleSheet(self._susp_style(checked))
        self._populate_table()

    # ─────────────────────────────────────────────────────────────────────────
    # Table card
    # ─────────────────────────────────────────────────────────────────────────

    def _build_table_card(self) -> QFrame:
        card = _card()
        card.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        vb = QVBoxLayout(card)
        vb.setContentsMargins(0, 0, 0, 0)
        vb.setSpacing(0)

        self.table = QTableWidget()
        self.table.verticalHeader().setVisible(False)
        self.table.setShowGrid(False)
        self.table.setAlternatingRowColors(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setFocusPolicy(Qt.NoFocus)
        self.table.setFrameShape(QFrame.NoFrame)
        self.table.setSortingEnabled(False)
        self.table.setWordWrap(False)
        self.table.verticalHeader().setDefaultSectionSize(44)

        self._status_delegate = _StatusDelegate(self.table)

        self.table.setStyleSheet(
            "QTableWidget{background:transparent;border:none;outline:none;}"
            "QTableWidget::item{"
            "  padding:6px 14px;border:none;"
            "  border-bottom:1px solid rgba(51,65,85,0.20);"
            "}"
            "QTableWidget::item:hover{background:rgba(51,65,85,0.28);}"
            "QTableWidget::item:selected{background:rgba(6,182,212,0.16);color:#ffffff;}"
            "QTableWidget::item:focus{outline:0;border:none;}"
        )
        self.table.horizontalHeader().setStyleSheet(
            "QHeaderView{background:rgba(22,33,52,0.95);border:none;}"
            "QHeaderView::section{"
            "  background:rgba(22,33,52,0.95);color:#94a3b8;"
            "  padding:11px 14px;border:none;"
            "  border-bottom:1px solid rgba(51,65,85,0.6);"
            "  font-size:11px;font-weight:700;letter-spacing:0.4px;"
            "}"
            "QHeaderView::section:hover{background:rgba(6,182,212,0.09);color:#e2e8f0;}"
        )
        self.table.horizontalHeader().sectionClicked.connect(self._header_clicked)

        sb_qss = (
            f"QScrollBar:vertical{{background:rgba(15,23,42,0.4);width:8px;"
            f"border-radius:4px;margin-top:{_HEADER_H}px;margin-bottom:0px;}}"
            "QScrollBar::handle:vertical{background:#334155;border-radius:4px;min-height:24px;}"
            "QScrollBar::handle:vertical:hover{background:#475569;}"
            "QScrollBar::add-line:vertical,QScrollBar::sub-line:vertical{height:0;background:transparent;border:none;}"
            "QScrollBar::add-page:vertical,QScrollBar::sub-page:vertical{background:rgba(15,23,42,0.4);border-radius:4px;}"
            "QScrollBar:horizontal{background:rgba(15,23,42,0.4);height:8px;border-radius:4px;margin:0px;}"
            "QScrollBar::handle:horizontal{background:#334155;border-radius:4px;min-width:24px;}"
            "QScrollBar::handle:horizontal:hover{background:#475569;}"
            "QScrollBar::add-line:horizontal,QScrollBar::sub-line:horizontal{width:0;background:transparent;border:none;}"
            "QScrollBar::add-page:horizontal,QScrollBar::sub-page:horizontal{background:rgba(15,23,42,0.4);border-radius:4px;}"
        )
        self.table.verticalScrollBar().setStyleSheet(sb_qss)
        self.table.horizontalScrollBar().setStyleSheet(sb_qss)

        self._empty_w = QWidget()
        self._empty_w.setStyleSheet("background:transparent;")
        ev = QVBoxLayout(self._empty_w)
        ev.setAlignment(Qt.AlignCenter)
        ev.setSpacing(10)
        ev.addWidget(_lbl("📂  No artifacts loaded", 16, True, "#334155"))
        ev.addWidget(_lbl("Run an analysis first to populate this view.", 12, color="#475569"))

        vb.addWidget(self.table, stretch=1)
        vb.addWidget(self._empty_w, stretch=1)
        self._empty_w.hide()
        return card

    # ─────────────────────────────────────────────────────────────────────────
    # Column setup
    # ─────────────────────────────────────────────────────────────────────────

    # Per-tab table widths tuned for readability. Long cells still expose
    # full values through tooltips.
    _COL_WIDTHS = {
        # Processes: ImageFileName column is export-only and hidden in the UI.
        0: [80, 1, 250, 80, 360, 90, 100, 155, 360],
        1: [80, 80, 260, 280, 170, 155, 360],
        # DLLs: Process column is export-only and hidden in the UI.
        2: [80, 1, 260, 390, 190, 95, 90, 90,
            86, 86, 86, 112, 112, 112, 150, 140, 155, 300],
        3: [145, 420, 80, 210, 155, 380],
        4: [220, 430, 200, 170, 170, 170, 95, 155, 360],
    }

    # Stretch one explanation/text-heavy column on wider screens.
    _STRETCH_COLS = {
        0: [8],
        1: [6],
        3: [5],
        4: [8],
    }

    def _reset_table(self, headers: list, narrow_cols: list, status_col: int):
        self.table.clearContents()
        self.table.setRowCount(0)
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)

        # Header tooltips help beginner users understand each recovered field.
        for c, header in enumerate(headers):
            item = self.table.horizontalHeaderItem(c)
            if item:
                clean_header = header.rstrip(" ▲▼")
                item.setToolTip(_HEADER_TIPS.get(clean_header, clean_header))

        for c in range(len(headers) + 1):
            self.table.setItemDelegateForColumn(c, None)
        self.table.setItemDelegateForColumn(status_col, self._status_delegate)

        hh = self.table.horizontalHeader()
        hh.setStretchLastSection(False)
        hh.setMinimumSectionSize(56)

        for c in range(len(headers)):
            hh.setSectionResizeMode(c, QHeaderView.Interactive)

        widths = self._COL_WIDTHS.get(self._active_tab, [])
        for c, w in enumerate(widths[:len(headers)]):
            self.table.setColumnWidth(c, w)

        if 0 <= status_col < len(headers):
            self.table.setColumnWidth(status_col, max(self.table.columnWidth(status_col), 155))

        for c in self._STRETCH_COLS.get(self._active_tab, []):
            if c < len(headers):
                hh.setSectionResizeMode(c, QHeaderView.Stretch)

        # Reset visibility; specific tab builders may hide export-only columns.
        for c in range(len(headers)):
            self.table.setColumnHidden(c, False)

        if 0 <= self._sort_col < len(headers):
            arrow = " ▲" if self._sort_asc else " ▼"
            hh.model().setHeaderData(
                self._sort_col, Qt.Horizontal,
                headers[self._sort_col] + arrow,
            )

    def _build_pid_name_map(self) -> dict:
        pid_map: dict = {}
        for p, _ in _flat_procs(AppState.current.processes):
            pid_str = _proc_pid(p)
            try:
                pid_map[int(pid_str)] = _proc_name(p)
            except (ValueError, TypeError):
                pass
        return pid_map

    # ─────────────────────────────────────────────────────────────────────────
    # Population dispatch
    # ─────────────────────────────────────────────────────────────────────────

    def _populate_table(self):
        {0: self._show_processes,
         1: self._show_threads,
         2: self._show_dlls,
         3: self._show_handles,
         4: self._show_pe_sections}[self._active_tab]()

    # ── Processes ─────────────────────────────────────────────────────────────

    def _show_processes(self):
        sess = AppState.current
        flat = _flat_procs(sess.processes)
        if self._susp_only:
            flat = [(p, d) for p, d in flat if _proc_status(p) == "suspicious"]

        t_counts = _count_by_pid(sess.threads)
        h_counts = _count_by_pid(sess.handles)

        def _tc(p) -> str:
            try: return str(t_counts.get(int(_proc_pid(p)), 0))
            except: return "0"
        def _hc(p) -> str:
            try: return str(h_counts.get(int(_proc_pid(p)), 0))
            except: return "0"

        getters = [
            lambda x: _proc_pid(x[0]),
            lambda x: _proc_image_file_name(x[0]),
            lambda x: _proc_name(x[0]),
            lambda x: _proc_ppid(x[0]),
            lambda x: _proc_path(x[0]),
            lambda x: _tc(x[0]),
            lambda x: _hc(x[0]),
            lambda x: _proc_status(x[0]),
            lambda x: _j10_conf_label(x[0]),
            lambda x: _proc_anomaly(x[0]),
        ]
        flat = _sorted_rows(flat, getters, self._sort_col, self._sort_asc)

        headers = ["PID", "ImageFileName", "Process Name", "PPID", "Path",
                   "Threads", "Handles", "Status", "Confidence", "Anomaly"]
        self._reset_table(headers, narrow_cols=[0, 1, 3, 5, 6], status_col=7)
        # Keep ImageFileName for CSV/Excel export, but hide it from the UI.
        self.table.setColumnHidden(1, True)
        self.table.setRowCount(len(flat))

        for row, (p, depth) in enumerate(flat):
            status  = _proc_status(p)
            is_susp = status == "suspicious"
            nc = "#f87171" if is_susp else "#e2e8f0"
            self._si(row, 0, _proc_pid(p),                    "#94a3b8", mono=True)
            self._si(row, 1, _proc_image_file_name(p),          "#94a3b8", mono=True)
            self._si(row, 2, "  "*depth+_proc_name(p),        nc, bold=is_susp)
            proc_item = self.table.item(row, 2)
            if proc_item:
                proc_item.setToolTip(_compact_tooltip(
                    f"Process Name: {_proc_name(p)}\n"
                    f"Image File Name: {_proc_image_file_name(p)}\n"
                    f"PID: {_proc_pid(p)}\nPPID: {_proc_ppid(p)}",
                    limit=300
                ))
            self._si(row, 3, _proc_ppid(p),                   "#64748b", mono=True)
            self._si(row, 4, _proc_path(p),                   "#64748b", mono=True, small=True, tip=True)
            self._si(row, 5, _tc(p),                          "#94a3b8")
            self._si(row, 6, _hc(p),                          "#94a3b8")
            self._status_item(row, 7, status)
            conf_label = _j10_conf_label(p)
            conf_col = "#fbbf24" if "Low-confidence" in conf_label else ("#93c5fd" if "Medium" in conf_label else "#34d399")
            self._si(row, 8, conf_label, conf_col, bold=True, small=True)
            if self.table.item(row, 8):
                self.table.item(row, 8).setToolTip(_compact_tooltip(_j10_conf_tooltip(p), limit=900))
            anom = _proc_anomaly(p)
            self._si(row, 9, anom, "#f87171" if anom!="-" else "#475569", bold=anom!="-")
            if anom != "-" and self.table.item(row, 9):
                self.table.item(row, 9).setToolTip(_anomaly_tooltip(anom))

        self._show_empty_if(not flat and not sess.processes)

    # ── Threads ───────────────────────────────────────────────────────────────

    def _show_threads(self):
        sess    = AppState.current
        pid_map = self._build_pid_name_map()
        data    = [t for t in sess.threads
                   if not self._susp_only or _thread_status(t) == "suspicious"]

        getters = [
            lambda x: _thread_tid(x),
            lambda x: _thread_pid(x),
            lambda x: _thread_process(x, pid_map),
            lambda x: _thread_start_addr(x),
            lambda x: _thread_state(x),
            lambda x: _thread_status(x),
            lambda x: _thread_anomaly(x),
        ]
        data = _sorted_rows(data, getters, self._sort_col, self._sort_asc)

        headers = ["TID", "PID", "Process", "Start Address", "State", "Status", "Anomaly"]
        self._reset_table(headers, narrow_cols=[0, 1, 2, 4], status_col=5)
        self.table.setRowCount(len(data))

        for row, t in enumerate(data):
            status  = _thread_status(t)
            is_susp = status == "suspicious"
            self._si(row, 0, _thread_tid(t),             "#94a3b8", mono=True)
            self._si(row, 1, _thread_pid(t),             "#64748b", mono=True)
            self._si(row, 2, _thread_process(t, pid_map),
                     "#f87171" if is_susp else "#e2e8f0")
            self._si(row, 3, _thread_start_addr(t),      "#06b6d4", mono=True)
            state_txt = _thread_state(t)
            self._si(row, 4, state_txt,                  "#94a3b8")
            if self.table.item(row, 4):
                self.table.item(row, 4).setToolTip(_compact_tooltip(
                    _thread_state_tooltip(state_txt), limit=360
                ))
            self._status_item(row, 5, status)
            anom = _thread_anomaly(t)
            self._si(row, 6, anom, "#f87171" if anom!="-" else "#475569", bold=anom!="-")
            if anom != "-" and self.table.item(row, 6):
                self.table.item(row, 6).setToolTip(_anomaly_tooltip(anom))

        self._show_empty_if(not data and not sess.threads)

    # ── DLLs ─────────────────────────────────────────────────────────────────

    def _show_dlls(self):
        sess = AppState.current
        pid_map = self._build_pid_name_map()

        def _dll_proc_display(d):
            direct = _dll_process(d)
            if direct not in ("-", "?"):
                return direct
            try:
                return pid_map.get(int(_dll_pid(d)), "-")
            except Exception:
                return "-"

        data = [d for d in sess.dlls
                if not self._susp_only or _dll_status(d) == "suspicious"]

        getters = [
            lambda x: _dll_pid(x),
            lambda x: _dll_proc_display(x),
            lambda x: _dll_name(x),
            lambda x: _dll_path(x),
            lambda x: _dll_base(x),
            lambda x: _dll_size(x),
            lambda x: "Yes" if _dll_signed(x) else "No",
            lambda x: str(_dll_entropy(x)),
            lambda x: _dll_bool_field(x, "in_load", "InLoad"),
            lambda x: _dll_bool_field(x, "in_mem", "in_memory", "InMemory"),
            lambda x: _dll_bool_field(x, "in_init", "InInit", "in_initialization"),
            lambda x: _dll_bool_field(x, "not_in_load", "NotInLoad"),
            lambda x: _dll_bool_field(x, "not_in_mem", "NotInMemory", "not_in_memory"),
            lambda x: _dll_bool_field(x, "not_in_init", "NotInInit", "not_in_initialization"),
            lambda x: _dll_bool_field(x, "ldr_lists_complete", "LdrListsComplete"),
            lambda x: _dll_text_field(x, "ldr_source", "LdrSource", default="-"),
            lambda x: _dll_status(x),
            None,
        ]
        data = _sorted_rows(data, getters, self._sort_col, self._sort_asc)

        headers = [
            "PID", "Process", "DLL Name", "Path", "Base Address", "Size",
            "Signed", "Entropy",
            "in_load", "in_mem", "in_init",
            "not_in_load", "not_in_mem", "not_in_init",
            "ldr_lists_complete", "ldr_source",
            "Status", "Anomaly",
        ]
        self._reset_table(headers, narrow_cols=[0, 1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14], status_col=16)
        # Keep Process for CSV/Excel export, but hide it from the UI.
        self.table.setColumnHidden(1, True)
        self.table.setRowCount(len(data))

        for row, d in enumerate(data):
            status  = _dll_status(d)
            is_susp = status == "suspicious"
            signed  = _dll_signed(d)
            ent     = _dll_entropy(d)
            proc_display = _dll_proc_display(d)
            self._si(row, 0, _dll_pid(d), "#64748b", mono=True)
            if self.table.item(row, 0):
                self.table.item(row, 0).setToolTip(_compact_tooltip(
                    f"PID: {_dll_pid(d)}\nProcess: {proc_display}", limit=300
                ))
            self._si(row, 1, proc_display, "#94a3b8")
            self._si(row, 2, _dll_name(d), "#f87171" if is_susp else "#e2e8f0", bold=is_susp)
            self._si(row, 3, _dll_path(d), "#64748b", mono=True, small=True, tip=True)
            self._si(row, 4, _dll_base(d), "#06b6d4", mono=True)
            self._si(row, 5, _dll_size(d), "#94a3b8")
            self._si(row, 6, "Yes" if signed else "No",
                     "#34d399" if signed else "#f87171", bold=not signed)
            self._si(row, 7, f"{ent:.1f}", "#f87171" if ent > 7.0 else "#94a3b8")

            # Native 3-list LDR membership columns.  Unknown/missing values are
            # displayed as '-' instead of false so old/parser-incomplete rows do
            # not look like suspicious mismatches.
            ldr_vals = [
                _dll_bool_field(d, "in_load", "InLoad"),
                _dll_bool_field(d, "in_mem", "in_memory", "InMemory"),
                _dll_bool_field(d, "in_init", "InInit", "in_initialization"),
                _dll_bool_field(d, "not_in_load", "NotInLoad"),
                _dll_bool_field(d, "not_in_mem", "NotInMemory", "not_in_memory"),
                _dll_bool_field(d, "not_in_init", "NotInInit", "not_in_initialization"),
                _dll_bool_field(d, "ldr_lists_complete", "LdrListsComplete"),
                _dll_text_field(d, "ldr_source", "LdrSource", default="-"),
            ]
            for idx, val in enumerate(ldr_vals, start=8):
                color = "#94a3b8"
                bold = False
                if idx in (11, 12, 13) and val == "Yes":
                    color = "#fbbf24"
                    bold = True
                elif val == "No":
                    color = "#64748b"
                elif val == "Yes":
                    color = "#34d399"
                self._si(row, idx, val, color, bold=bold)

            self._status_item(row, 16, status)
            anoms = d.get("anomalies", [])
            anom_str = ("; ".join(
                a.get("type", str(a)) if isinstance(a, dict) else str(a)
                for a in anoms[:2]) if anoms else d.get("anomaly", "-") or "-")
            self._si(row, 17, anom_str,
                     "#f87171" if anom_str!="-" else "#475569", bold=anom_str!="-")
            if anom_str != "-" and self.table.item(row, 17):
                self.table.item(row, 17).setToolTip(_anomaly_tooltip(anom_str))

        self._show_empty_if(not data and not sess.dlls)

    # ── Handles ───────────────────────────────────────────────────────────────

    def _show_handles(self):
        sess    = AppState.current
        pid_map = self._build_pid_name_map()

        # Pre-compute status+anomaly for every handle (mirrors tool heuristics)
        computed = [_handle_compute(h) for h in sess.handles]

        data = [(h, st, an) for h, (st, an) in zip(sess.handles, computed)
                if not self._susp_only or st == "suspicious"]

        getters = [
            lambda x: _handle_type(x[0]),
            lambda x: _handle_name(x[0]),
            lambda x: _handle_pid(x[0]),
            lambda x: _handle_process(x[0], pid_map),
            lambda x: x[1],   # status
            None,
        ]
        data = _sorted_rows(data, getters, self._sort_col, self._sort_asc)

        headers = ["Type", "Handle Name", "PID", "Process", "Status", "Anomaly"]
        self._reset_table(headers, narrow_cols=[0, 2, 3], status_col=4)
        self.table.setRowCount(len(data))

        for row, (h, status, anom_str) in enumerate(data):
            is_susp = status == "suspicious"
            self._si(row, 0, _handle_type(h),             "#06b6d4")
            self._si(row, 1, _handle_name(h),
                     "#f87171" if is_susp else "#e2e8f0",
                     mono=True, small=True, tip=True)
            self._si(row, 2, _handle_pid(h),              "#94a3b8", mono=True)
            self._si(row, 3, _handle_process(h, pid_map), "#f87171" if is_susp else "#e2e8f0")
            self._status_item(row, 4, status)
            self._si(row, 5, anom_str,
                     "#f87171" if anom_str!="-" else "#475569", bold=anom_str!="-")
            if anom_str != "-" and self.table.item(row, 5):
                self.table.item(row, 5).setToolTip(_anomaly_tooltip(anom_str))

        self._show_empty_if(not data and not sess.handles)

    # ── PE Sections ───────────────────────────────────────────────────────────

    def _show_pe_sections(self):
        sess = AppState.current
        pid_map = self._build_pid_name_map() 
        
        # --- FIX 1: Merge Injections into PE Sections ---
        raw_sections = list(sess.pe_sections)
        if hasattr(sess, "injections") and sess.injections:
            for inj in sess.injections:
                inj["is_suspicious"] = True  # Force it to highlight red!
                if "anomaly" not in inj and "anomalies" not in inj:
                    inj["anomaly"] = "Code Injection"
            raw_sections.extend(sess.injections)

        data = [s for s in raw_sections
                if not self._susp_only or _pe_status(s) == "suspicious"]

        # --- FIX 2: Add the VAD/Injection keys so data populates ---
        getters = [
            lambda x: _pe_proc(x, pid_map),
            lambda x: _pe_section(x),
            lambda x: _pe_virtual_addr_display(x),
            lambda x: _pe_virtual_size_display(x),
            lambda x: _pe_raw_size_display(x),
            lambda x: _pe_perms(x),
            lambda x: _pe_entropy_display(x),
            lambda x: _pe_status(x),
            None,
        ]
        data = _sorted_rows(data, getters, self._sort_col, self._sort_asc)

        headers = ["Process", "Section", "Virtual Addr", "Virtual Size",
                   "Raw Size", "Permissions", "Entropy", "Status", "Anomaly"]
        self._reset_table(headers, narrow_cols=[1,2,3,4,5,6], status_col=7)
        self.table.setRowCount(len(data))

        for row, s in enumerate(data):
            status  = _pe_status(s)
            is_susp = status == "suspicious"
            ent     = _pe_entropy(s)
            perms   = _pe_perms(s)
            rwx     = "W" in str(perms) and "X" in str(perms)
            
            va_display = _pe_virtual_addr_display(s)
            vs_display = _pe_virtual_size_display(s)
            raw_display = _pe_raw_size_display(s)
            entropy_display = _pe_entropy_display(s)

            self._si(row, 0, _pe_proc(s, pid_map),   "#f87171" if is_susp else "#e2e8f0")
            self._si(row, 1, _pe_section(s), "#e2e8f0", mono=True)
            self._si(row, 2, va_display, "#94a3b8", mono=True)
            self._si(row, 3, vs_display, "#94a3b8", mono=True)
            self._si(row, 4, raw_display, "#94a3b8", mono=True)

            for col, value, field_name in (
                (2, va_display, "Virtual Addr"),
                (3, vs_display, "Virtual Size"),
                (4, raw_display, "Raw Size"),
            ):
                if value in {"N/A", "?"} and self.table.item(row, col):
                    self.table.item(row, col).setToolTip(_compact_tooltip(
                        _pe_na_tooltip(field_name, s), limit=520
                    ))

            self._si(row, 5, perms, "#f87171" if rwx else "#94a3b8", bold=rwx)
            if self.table.item(row, 5):
                self.table.item(row, 5).setToolTip(_compact_tooltip(
                    _permission_tooltip(perms), limit=520
                ))

            self._si(row, 6, entropy_display, "#f87171" if (entropy_display != "N/A" and ent > 7.0) else "#94a3b8")
            if entropy_display == "N/A" and self.table.item(row, 6):
                self.table.item(row, 6).setToolTip(_compact_tooltip(
                    _pe_na_tooltip("Entropy", s),
                    limit=560
                ))
            self._status_item(row, 7, status)
            
            anom = s.get("anomaly", s.get("anomalies", "-") or "-")
            if isinstance(anom, list):
                anom = "; ".join(str(a) for a in anom[:2]) or "-"
            self._si(row, 8, str(anom),
                     "#f87171" if str(anom)!="-" else "#475569", bold=str(anom)!="-")
            if str(anom) != "-" and self.table.item(row, 8):
                self.table.item(row, 8).setToolTip(_anomaly_tooltip(str(anom)))

        self._show_empty_if(not data)

    # ─────────────────────────────────────────────────────────────────────────
    # Cell helpers
    # ─────────────────────────────────────────────────────────────────────────

    _F_NORMAL = QFont("Segoe UI", 12)
    _F_MONO   = QFont("Courier New", 12)
    _F_SMALL  = QFont("Courier New", 10)

    def _si(self, row:int, col:int, text:str,
            color:str="#e2e8f0", mono:bool=False,
            bold:bool=False, small:bool=False, tip:bool=False):
        item = QTableWidgetItem(str(text))
        base = self._F_SMALL if (small and mono) else (self._F_MONO if mono else self._F_NORMAL)
        f = QFont(base); f.setBold(bold)
        item.setFont(f)
        item.setForeground(QColor(color))
        if tip and str(text).strip():
            item.setToolTip(_compact_tooltip(str(text)))
        self.table.setItem(row, col, item)

    def _status_item(self, row:int, col:int, status:str):
        is_susp = status == "suspicious"
        item = QTableWidgetItem("⚠  SUSPICIOUS" if is_susp else "✓  CLEAN")
        item.setTextAlignment(Qt.AlignCenter)
        if is_susp:
            item.setToolTip(
                "Definition: SUSPICIOUS means this artifact has one or more "
                "artifact-level anomaly indicators for analyst review."
            )
        item.setBackground(QBrush(QColor(0, 0, 0, 0)))
        self.table.setItem(row, col, item)

    def _show_empty_if(self, condition:bool):
        if condition:
            self.table.hide(); self._empty_w.show()
        else:
            self._empty_w.hide(); self.table.show()


    def _current_export_verdict(self) -> str | None:
        """Best-effort verdict lookup for audit logs only."""
        try:
            sess = AppState.current
            for attr in ("verdict", "label"):
                val = getattr(sess, attr, None)
                if val:
                    return str(val).upper()
            for attr in ("classification_result", "result", "ml_result"):
                obj = getattr(sess, attr, None)
                if isinstance(obj, dict):
                    val = obj.get("verdict") or obj.get("label")
                    if val:
                        return str(val).upper()
        except Exception:
            pass
        return None

    def _audit_artifact_export(self, path: str, fmt: str, tab: str, *, rows: int | None = None, all_tabs: bool = False) -> None:
        """Register Artifact Viewer exports in Activity Monitor.

        This is monitoring-only. It never touches parser output, FeatureMapper,
        Autoencoder MSE fusion, RF threshold, or the final ML verdict.
        """
        try:
            from ramdetect.case_management.activity_logger import ActivityLogger, Action

            sess = AppState.current
            dump_name = getattr(sess, "dump_filename", "") if sess else ""
            fmt_norm = str(fmt or "").upper()
            if fmt_norm == "CSV":
                action = Action.CSV_EXPORTED
            elif all_tabs:
                action = Action.WORKBOOK_EXPORTED
            else:
                action = Action.XLSX_EXPORTED

            ActivityLogger.instance().log(
                action,
                file_path=path,
                file_name=os.path.basename(path),
                report_path=path,
                verdict=self._current_export_verdict(),
                details={
                    "source": "artifacts_viewer",
                    "format": fmt_norm,
                    "tab": tab,
                    "all_tabs": bool(all_tabs),
                    "rows": rows,
                    "dump": dump_name,
                },
            )
        except Exception:
            # Export must not fail because monitoring is unavailable.
            pass


    # ─────────────────────────────────────────────────────────────────────────
    # Search / sort / export
    # ─────────────────────────────────────────────────────────────────────────

    def _apply_row_filter(self, text:str=""):
        text = text.lower().strip()
        for r in range(self.table.rowCount()):
            if not text:
                self.table.setRowHidden(r, False); continue
            match = any(
                self.table.item(r, c) and
                text in self.table.item(r, c).text().lower()
                for c in range(self.table.columnCount())
            )
            self.table.setRowHidden(r, not match)

    def _header_clicked(self, col:int):
        if self._sort_col == col:
            self._sort_asc = not self._sort_asc
        else:
            self._sort_col = col; self._sort_asc = True
        self._populate_table()
        if self.search.text():
            self._apply_row_filter(self.search.text())

    def _export_csv(self):
        tab  = TABS[self._active_tab]
        ts   = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"{tab.lower().replace(' ','_')}_{ts}.csv"

        # ── Per-dump subfolder ─────────────────────────────────────────────
        # Save into reports/<dump_name>/ so exports stay organised.
        outdir = self._dump_export_dir()
        path = os.path.join(outdir, fname)

        # Status cell text → plain CSV text (strip emoji badges)
        _status_clean = {
            "⚠  SUSPICIOUS": "SUSPICIOUS",
            "✓  CLEAN":      "CLEAN",
            "⚠ SUSPICIOUS":  "SUSPICIOUS",
            "✓ CLEAN":       "CLEAN",
        }

        # Snapshot filter state, then rebuild table with everything visible
        prev_susp   = self._susp_only
        prev_search = self.search.text()
        self._susp_only = False
        self.search.blockSignals(True)
        self.search.clear()
        self._populate_table()

        headers = [self.table.horizontalHeaderItem(c).text().rstrip(" ▲▼")
                   for c in range(self.table.columnCount())]
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                # Standard CSV quoting keeps Windows paths readable:
                # C:\Windows\System32 instead of C:\\Windows\\System32.
                w = csv.writer(f, quoting=csv.QUOTE_ALL)
                w.writerow([_csv_safe_cell(h) for h in headers])
                for r in range(self.table.rowCount()):   # all rows, no filter
                    row_data = []
                    for c in range(self.table.columnCount()):
                        cell = self.table.item(r, c)
                        text = cell.text() if cell else ""
                        text = _status_clean.get(text, text)
                        row_data.append(_csv_safe_cell(text))
                    w.writerow(row_data)

            size = os.path.getsize(path)
            sz   = f"{size/1024:.1f} KB" if size >= 1024 else f"{size} B"
            toast = _ToastDialog(
                "Export successful",
                f"{fname}\n{sz}  —  {outdir}",
                ms=4000,
                parent=self,
            )
            toast.show_near(self)
            self._audit_artifact_export(path, "CSV", tab, rows=self.table.rowCount(), all_tabs=False)

        except Exception as e:
            toast = _ToastDialog("Export failed", str(e), ms=6000, parent=self)
            toast.show_near(self)

        finally:
            # Restore previous filter state
            self._susp_only = prev_susp
            self.search.setText(prev_search)
            self.search.blockSignals(False)
            self._populate_table()

    def _dump_export_dir(self) -> str:
        """Return reports/<dump_name>/ and create it if needed."""
        import re
        sess = AppState.current
        raw_name = (sess.dump_filename if sess and sess.dump_filename else "") or ""
        stem = os.path.splitext(raw_name)[0].strip()
        stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem)
        dump_sub = stem or "unknown_dump"
        outdir = os.path.join(_project_reports_dir(), dump_sub)
        os.makedirs(outdir, exist_ok=True)
        return outdir

    def _table_snapshot(self) -> tuple:
        """Return current table headers and all rows as plain text."""
        _status_clean = {
            "⚠  SUSPICIOUS": "SUSPICIOUS",
            "✓  CLEAN":      "CLEAN",
            "⚠ SUSPICIOUS":  "SUSPICIOUS",
            "✓ CLEAN":       "CLEAN",
        }
        headers = [
            self.table.horizontalHeaderItem(c).text().rstrip(" ▲▼")
            for c in range(self.table.columnCount())
        ]
        rows = []
        for r in range(self.table.rowCount()):
            row_data = []
            for c in range(self.table.columnCount()):
                cell = self.table.item(r, c)
                value = cell.text() if cell else ""
                row_data.append(_status_clean.get(value, value))
            rows.append(row_data)
        return headers, rows

    def _style_artifact_worksheet(self, ws, headers):
        """Apply RAMDetect-style formatting to an artifact worksheet."""
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="162134")
        clean_fill = PatternFill("solid", fgColor="D1FAE5")
        susp_fill = PatternFill("solid", fgColor="FEE2E2")
        thin = Side(style="thin", color="CBD5E1")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        status_col = None
        for i, h in enumerate(headers, start=1):
            if h == "Status":
                status_col = i
                break

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)
            if status_col:
                st = str(row[status_col - 1].value or "").upper()
                if st == "SUSPICIOUS":
                    row[status_col - 1].fill = susp_fill
                    row[status_col - 1].font = Font(bold=True, color="B91C1C")
                elif st == "CLEAN":
                    row[status_col - 1].fill = clean_fill
                    row[status_col - 1].font = Font(bold=True, color="047857")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for col_idx, header in enumerate(headers, start=1):
            max_len = len(str(header))
            # Use first 80 rows for width estimation to avoid slow exports.
            for row_idx in range(2, min(ws.max_row, 80) + 1):
                v = ws.cell(row_idx, col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            width = min(max(max_len + 2, 10), 38)
            if header in ("Path", "Handle Name", "Anomaly", "Section"):
                width = min(max(width, 24), 45)
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _show_export_choice_dialog(self):
        """Small RAMDetect-themed chooser for current-tab export format."""
        current_tab = TABS[self._active_tab]

        dlg = QDialog(self)
        dlg.setModal(True)
        dlg.setWindowTitle("Export current artifact tab")
        dlg.setWindowFlags(Qt.Dialog | Qt.FramelessWindowHint)
        dlg.setAttribute(Qt.WA_TranslucentBackground, True)
        dlg.setFixedSize(420, 230)

        outer = QVBoxLayout(dlg)
        outer.setContentsMargins(10, 10, 10, 10)

        card = QFrame()
        card.setObjectName("exportChoiceCard")
        card.setStyleSheet(
            "QFrame#exportChoiceCard {"
            " background: #0f172a;"
            " border: 1px solid rgba(6,182,212,0.38);"
            " border-radius: 14px;"
            "}"
            "QLabel { background: transparent; border: none; }"
        )
        outer.addWidget(card)

        root = QVBoxLayout(card)
        root.setContentsMargins(20, 16, 20, 18)
        root.setSpacing(14)

        # Header
        header = QHBoxLayout()
        header.setSpacing(8)

        title = QLabel("Export Current Tab")
        title.setFont(QFont("Segoe UI", 14, QFont.Bold))
        title.setStyleSheet("color: #ffffff;")
        header.addWidget(title, 1)

        close_btn = QPushButton("×")
        close_btn.setCursor(Qt.PointingHandCursor)
        close_btn.setFixedSize(28, 28)
        close_btn.setStyleSheet(
            "QPushButton {"
            " background: transparent; color: #cbd5e1; border: none;"
            " border-radius: 14px; font-size: 20px; font-weight: 700; padding: 0px;"
            "}"
            "QPushButton:hover { background: rgba(239,68,68,0.18); color: white; }"
        )
        close_btn.clicked.connect(dlg.reject)
        header.addWidget(close_btn)
        root.addLayout(header)

        subtitle = QLabel(f"Choose an export format for: {current_tab}")
        subtitle.setFont(QFont("Segoe UI", 10))
        subtitle.setStyleSheet("color: #94a3b8;")
        root.addWidget(subtitle)

        # Choices
        choices = QHBoxLayout()
        choices.setSpacing(12)

        csv_btn = QPushButton("⬇  CSV\nRaw table data")
        csv_btn.setCursor(Qt.PointingHandCursor)
        csv_btn.setFixedSize(178, 78)
        csv_btn.setStyleSheet(
            "QPushButton {"
            " background: rgba(30,41,59,0.72);"
            " color: #e2e8f0;"
            " border: 1px solid rgba(51,65,85,0.75);"
            " border-radius: 12px;"
            " font-size: 11px;"
            " font-weight: 700;"
            " text-align: center;"
            "}"
            "QPushButton:hover {"
            " background: rgba(6,182,212,0.20);"
            " border: 1px solid rgba(103,232,249,0.65);"
            " color: #67e8f9;"
            "}"
            "QPushButton:pressed {"
            " background: rgba(8,145,178,0.28);"
            " border: 1px solid rgba(6,182,212,0.80);"
            " color: #ffffff;"
            "}"
        )
        choices.addWidget(csv_btn)

        xlsx_btn = QPushButton("📄  XLSX\nThemed workbook")
        xlsx_btn.setCursor(Qt.PointingHandCursor)
        xlsx_btn.setFixedSize(178, 78)
        xlsx_btn.setStyleSheet(
            "QPushButton {"
            " background: rgba(30,41,59,0.72);"
            " color: #e2e8f0;"
            " border: 1px solid rgba(51,65,85,0.75);"
            " border-radius: 12px;"
            " font-size: 11px;"
            " font-weight: 700;"
            " text-align: center;"
            "}"
            "QPushButton:hover {"
            " background: rgba(6,182,212,0.20);"
            " border: 1px solid rgba(103,232,249,0.65);"
            " color: #67e8f9;"
            "}"
            "QPushButton:pressed {"
            " background: rgba(8,145,178,0.28);"
            " border: 1px solid rgba(6,182,212,0.80);"
            " color: #ffffff;"
            "}"
        )
        choices.addWidget(xlsx_btn)
        root.addLayout(choices)

        note = QLabel("CSV is lightweight. XLSX keeps the RAMDetect report theme.")
        note.setFont(QFont("Segoe UI", 9))
        note.setStyleSheet("color: #64748b;")
        note.setAlignment(Qt.AlignCenter)
        root.addWidget(note)

        def do_csv():
            dlg.accept()
            self._export_csv()

        def do_xlsx():
            dlg.accept()
            self._export_excel_current()

        csv_btn.clicked.connect(do_csv)
        xlsx_btn.clicked.connect(do_xlsx)

        try:
            parent_geo = self.window().frameGeometry()
            dlg.move(parent_geo.center() - dlg.rect().center())
        except Exception:
            pass

        dlg.exec_()

    def _export_excel_current(self):
        """Export the current artifact tab as one themed XLSX workbook."""
        try:
            from openpyxl import Workbook
        except Exception:
            toast = _ToastDialog(
                "Export failed",
                "openpyxl is required for Excel export. Install it with: pip install openpyxl",
                ms=7000,
                parent=self,
            )
            toast.show_near(self)
            return

        tab = TABS[self._active_tab]
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"{tab.lower().replace(' ','_')}_{ts}.xlsx"
        outdir = self._dump_export_dir()
        path = os.path.join(outdir, fname)

        prev_susp = self._susp_only
        prev_search = self.search.text()
        prev_sort_col = self._sort_col
        prev_sort_asc = self._sort_asc

        try:
            self.search.blockSignals(True)
            self.search.clear()
            self._susp_only = False
            self._sort_col = -1
            self._sort_asc = True
            self._populate_table()

            headers, rows = self._table_snapshot()

            wb = Workbook()
            ws = wb.active
            ws.title = tab[:31]
            ws.append([_excel_safe_cell(v) for v in headers])
            for row_data in rows:
                ws.append([_excel_safe_cell(v) for v in row_data])

            self._style_artifact_worksheet(ws, headers)
            wb.save(path)

            size = os.path.getsize(path)
            sz = f"{size/1024:.1f} KB" if size >= 1024 else f"{size} B"
            toast = _ToastDialog(
                "Excel tab export successful",
                f"{fname}\n{len(rows)} rows · {sz}  —  {outdir}",
                ms=5000,
                parent=self,
            )
            toast.show_near(self)
            self._audit_artifact_export(path, "XLSX", tab, rows=len(rows), all_tabs=False)

        except Exception as e:
            toast = _ToastDialog("Export failed", str(e), ms=7000, parent=self)
            toast.show_near(self)

        finally:
            self._susp_only = prev_susp
            self.search.setText(prev_search)
            self.search.blockSignals(False)
            self._sort_col = prev_sort_col
            self._sort_asc = prev_sort_asc
            self._populate_table()
            if prev_search:
                self._apply_row_filter(prev_search)

    def _export_excel_all(self):
        """Export all artifact categories into one XLSX workbook.

        Each artifact tab becomes a separate Excel sheet:
        Processes, Threads, DLLs, Handles, PE Sections.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
            from openpyxl.utils import get_column_letter
        except Exception:
            toast = _ToastDialog(
                "Export failed",
                "openpyxl is required for Excel export. Install it with: pip install openpyxl",
                ms=7000,
                parent=self,
            )
            toast.show_near(self)
            return

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        fname = f"ramdetect_artifacts_all_{ts}.xlsx"
        outdir = self._dump_export_dir()
        path = os.path.join(outdir, fname)

        prev_tab = self._active_tab
        prev_sort_col = self._sort_col
        prev_sort_asc = self._sort_asc
        prev_susp = self._susp_only
        prev_search = self.search.text()

        wb = Workbook()
        default_ws = wb.active
        wb.remove(default_ws)
        total_rows = 0

        try:
            self.search.blockSignals(True)
            self.search.clear()
            self._susp_only = False
            self._sort_col = -1
            self._sort_asc = True

            for idx, tab in enumerate(TABS):
                self._active_tab = idx
                self._populate_table()
                headers, rows = self._table_snapshot()
                total_rows += len(rows)

                ws = wb.create_sheet(tab[:31])
                ws.append([_excel_safe_cell(v) for v in headers])
                for row_data in rows:
                    ws.append([_excel_safe_cell(v) for v in row_data])

                self._style_artifact_worksheet(ws, headers)

            wb.save(path)

            size = os.path.getsize(path)
            sz = f"{size/1024:.1f} KB" if size >= 1024 else f"{size} B"
            toast = _ToastDialog(
                "Excel export successful",
                f"{fname}\n5 sheets exported · {sz}  —  {outdir}",
                ms=5000,
                parent=self,
            )
            toast.show_near(self)
            self._audit_artifact_export(path, "XLSX", "All Artifacts", rows=total_rows, all_tabs=True)

        except Exception as e:
            toast = _ToastDialog("Export failed", str(e), ms=7000, parent=self)
            toast.show_near(self)

        finally:
            self._active_tab = prev_tab
            self._sort_col = prev_sort_col
            self._sort_asc = prev_sort_asc
            self._susp_only = prev_susp
            self.search.setText(prev_search)
            self.search.blockSignals(False)
            self._refresh_tab_styles()
            self._susp_btn.setChecked(self._susp_only)
            self._susp_btn.setStyleSheet(self._susp_style(self._susp_only))
            self._populate_table()
            if prev_search:
                self._apply_row_filter(prev_search)


    def showEvent(self, event):
        super().showEvent(event)
        self._refresh_tab_styles()
        self._populate_table()